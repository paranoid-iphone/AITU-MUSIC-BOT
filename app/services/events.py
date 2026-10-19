from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.timezone import now_local
from app.db.models import BandMember, Event, EventApplication, EventApplicationMember, EventStatus, PersonalEventHistory
from app.services.errors import ConflictError, ForbiddenError, NotFoundError
from app.services.users import get_user_by_telegram_id, require_admin, require_approved


def _request_sheet_sync() -> None:
    """The concerts sheet follows every committed change of a concert or its applications."""
    from app.services.sheets_sync import request_sync  # lazy: sheets_sync imports booking code

    request_sync()


async def create_event(session: AsyncSession, telegram_id: int, **data) -> Event:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    require_admin(user)
    event = Event(created_by_id=user.id, **data)
    session.add(event)
    await session.flush()
    await log_action(
        session,
        user.id,
        "event.create",
        "event",
        event.id,
        f"Создал концерт: {event.title} ({event.event_date.isoformat()})",
    )
    await session.commit()
    await session.refresh(event)
    _request_sheet_sync()
    return event


async def update_event(session: AsyncSession, event_id: int, telegram_id: int, **data) -> Event:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_admin(user)
    result = await session.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()
    if event is None:
        raise NotFoundError("Event not found")
    for key, value in data.items():
        setattr(event, key, value)
    changed = ", ".join(sorted(data.keys())) or "без изменений"
    await log_action(
        session,
        user.id,
        "event.update",
        "event",
        event.id,
        f"Обновил концерт #{event.id}: {changed}",
    )
    await session.commit()
    await session.refresh(event)
    _request_sheet_sync()
    return event


async def list_events(session: AsyncSession) -> list[Event]:
    result = await session.execute(select(Event).order_by(Event.event_date.desc()))
    return list(result.scalars())


async def _require_open_event(session: AsyncSession, event_id: int) -> Event:
    event = (await session.execute(select(Event).where(Event.id == event_id))).scalar_one_or_none()
    if event is None:
        raise NotFoundError("Event not found")
    if event.status != EventStatus.OPEN.value:
        raise ConflictError("Applications for this event are closed")
    if event.submission_deadline is not None and now_local().date() > event.submission_deadline:
        raise ConflictError("The submission deadline for this event has passed")
    return event


async def submit_application(
    session: AsyncSession,
    event_id: int,
    telegram_id: int,
    band_id: int,
    song_title: str,
    member_ids: list[int],
) -> EventApplication:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    owner_result = await session.execute(
        select(BandMember).where(BandMember.band_id == band_id, BandMember.user_id == user.id, BandMember.is_owner.is_(True))
    )
    if owner_result.scalar_one_or_none() is None:
        raise ForbiddenError("Only band owner can submit lists")
    await _require_open_event(session, event_id)
    members_result = await session.execute(
        select(BandMember).where(BandMember.band_id == band_id, BandMember.id.in_(member_ids)).options(selectinload(BandMember.user))
    )
    members = list(members_result.scalars())
    if not members:
        raise ConflictError("Choose at least one member")
    incomplete = [m.user for m in members if not m.user.profile_completed]
    if incomplete:
        names = ", ".join(user.telegram_username or str(user.telegram_id) for user in incomplete)
        raise ConflictError(f"Some members have incomplete profiles: {names}")
    existing = await session.execute(select(EventApplication).where(EventApplication.event_id == event_id, EventApplication.band_id == band_id))
    application = existing.scalar_one_or_none()
    if application is None:
        application = EventApplication(event_id=event_id, band_id=band_id, submitted_by_id=user.id, song_title=song_title)
        session.add(application)
        await session.flush()
    else:
        application.song_title = song_title
        application.submitted_by_id = user.id
        old_members = await session.execute(select(EventApplicationMember).where(EventApplicationMember.application_id == application.id))
        for old_member in old_members.scalars():
            await session.delete(old_member)
        await session.flush()  # the unit of work inserts before it deletes; a kept member would hit uq_application_user
    for member in members:
        session.add(EventApplicationMember(application_id=application.id, user_id=member.user_id, instrument_role=member.instrument_role))
    await log_action(
        session,
        user.id,
        "event.application.submit",
        "event_application",
        application.id,
        f"Подал номер на концерт #{event_id}: {song_title}",
    )
    await session.commit()
    await session.refresh(application)
    _request_sheet_sync()
    return application


async def list_event_applications(session: AsyncSession, telegram_id: int, event_id: int) -> list[dict]:
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_admin(user)
    result = await session.execute(
        select(EventApplication)
        .where(EventApplication.event_id == event_id)
        .options(
            selectinload(EventApplication.band),
            selectinload(EventApplication.members).selectinload(EventApplicationMember.user),
        )
        .order_by(EventApplication.submitted_at.desc())
    )
    rows = []
    for application in result.scalars().unique():
        rows.append(
            {
                "id": application.id,
                "event_id": application.event_id,
                "band_id": application.band_id,
                "band_name": application.band.name,
                "song_title": application.song_title,
                "submitted_at": application.submitted_at.isoformat(),
                "members": [
                    {
                        "id": member.id,
                        "instrument_role": member.instrument_role,
                        "user": {
                            "telegram_id": member.user.telegram_id,
                            "telegram_username": member.user.telegram_username,
                            "first_name": member.user.first_name,
                            "last_name": member.user.last_name,
                            "study_group": member.user.study_group,
                        },
                    }
                    for member in application.members
                ],
            }
        )
    return rows


def _application_to_row(application: EventApplication) -> dict:
    return {
        "id": application.id,
        "event_id": application.event_id,
        "event_title": application.event.title,
        "event_date": application.event.event_date.isoformat(),
        "event_location": application.event.location,
        "event_status": application.event.status,
        "band_id": application.band_id,
        "band_name": application.band.name,
        "song_title": application.song_title,
        "submitted_at": application.submitted_at.isoformat(),
        "members": [
            {
                "id": member.id,
                "instrument_role": member.instrument_role,
                "user": {
                    "telegram_id": member.user.telegram_id,
                    "telegram_username": member.user.telegram_username,
                    "first_name": member.user.first_name,
                    "last_name": member.user.last_name,
                    "study_group": member.user.study_group,
                },
            }
            for member in application.members
        ],
    }


async def _require_application_manager(session: AsyncSession, telegram_id: int, application: EventApplication):
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    owner_result = await session.execute(
        select(BandMember).where(
            BandMember.band_id == application.band_id,
            BandMember.user_id == user.id,
            BandMember.is_owner.is_(True),
        )
    )
    if owner_result.scalar_one_or_none() is None:
        raise ForbiddenError("Only band owner can manage this application")
    return user


async def list_my_applications(session: AsyncSession, telegram_id: int) -> list[dict]:
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    result = await session.execute(
        select(EventApplication)
        .join(BandMember, BandMember.band_id == EventApplication.band_id)
        .where(BandMember.user_id == user.id, BandMember.is_owner.is_(True))
        .options(
            selectinload(EventApplication.event),
            selectinload(EventApplication.band),
            selectinload(EventApplication.members).selectinload(EventApplicationMember.user),
        )
        .order_by(EventApplication.submitted_at.desc())
    )
    return [_application_to_row(application) for application in result.scalars().unique()]


async def get_application_detail(session: AsyncSession, telegram_id: int, application_id: int) -> dict:
    result = await session.execute(
        select(EventApplication)
        .where(EventApplication.id == application_id)
        .options(
            selectinload(EventApplication.event),
            selectinload(EventApplication.band),
            selectinload(EventApplication.members).selectinload(EventApplicationMember.user),
        )
    )
    application = result.scalar_one_or_none()
    if application is None:
        raise NotFoundError("Application not found")
    await _require_application_manager(session, telegram_id, application)
    return _application_to_row(application)


async def update_application(
    session: AsyncSession,
    application_id: int,
    telegram_id: int,
    song_title: str | None = None,
    member_ids: list[int] | None = None,
) -> EventApplication:
    from app.services.audit import log_action

    result = await session.execute(
        select(EventApplication)
        .where(EventApplication.id == application_id)
        .options(selectinload(EventApplication.event), selectinload(EventApplication.band))
    )
    application = result.scalar_one_or_none()
    if application is None:
        raise NotFoundError("Application not found")
    user = await _require_application_manager(session, telegram_id, application)
    await _require_open_event(session, application.event_id)
    if song_title is not None:
        application.song_title = song_title
    if member_ids is not None:
        members_result = await session.execute(
            select(BandMember)
            .where(BandMember.band_id == application.band_id, BandMember.id.in_(member_ids))
            .options(selectinload(BandMember.user))
        )
        members = list(members_result.scalars())
        if not members:
            raise ConflictError("Choose at least one member")
        incomplete = [m.user for m in members if not m.user.profile_completed]
        if incomplete:
            names = ", ".join(user.telegram_username or str(user.telegram_id) for user in incomplete)
            raise ConflictError(f"Some members have incomplete profiles: {names}")
        old_members = await session.execute(select(EventApplicationMember).where(EventApplicationMember.application_id == application.id))
        for old_member in old_members.scalars():
            await session.delete(old_member)
        await session.flush()  # the unit of work inserts before it deletes; a kept member would hit uq_application_user
        for member in members:
            session.add(EventApplicationMember(application_id=application.id, user_id=member.user_id, instrument_role=member.instrument_role))
    await log_action(
        session,
        user.id,
        "event.application.update",
        "event_application",
        application.id,
        f"Обновил заявку #{application.id}: {application.song_title or '-'}",
    )
    await session.commit()
    await session.refresh(application)
    _request_sheet_sync()
    return application


async def delete_application(session: AsyncSession, telegram_id: int, application_id: int) -> dict:
    from app.services.audit import log_action

    result = await session.execute(
        select(EventApplication)
        .where(EventApplication.id == application_id)
        .options(selectinload(EventApplication.event), selectinload(EventApplication.band))
    )
    application = result.scalar_one_or_none()
    if application is None:
        raise NotFoundError("Application not found")
    user = await _require_application_manager(session, telegram_id, application)
    deleted = {
        "id": application.id,
        "event_title": application.event.title,
        "band_name": application.band.name,
        "song_title": application.song_title,
    }
    await log_action(
        session,
        user.id,
        "event.application.delete",
        "event_application",
        application.id,
        f"Отозвал заявку #{application.id}: {application.song_title or '-'}",
    )
    await session.delete(application)
    await session.commit()
    _request_sheet_sync()
    return deleted


def _personal_event_to_row(item: PersonalEventHistory) -> dict:
    return {
        "id": item.id,
        "source": "personal",
        "source_label": "ты добавил",
        "event_title": item.title,
        "event_date": item.event_date.isoformat(),
        "event_location": item.location,
        "event_status": None,
        "band_id": None,
        "band_name": None,
        "song_title": None,
        "instrument_role": item.role_at_event,
        "created_at": item.created_at.isoformat(),
        "updated_at": item.updated_at.isoformat(),
    }


async def list_user_event_history(session: AsyncSession, telegram_id: int) -> list[dict]:
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    result = await session.execute(
        select(EventApplicationMember)
        .join(EventApplication, EventApplicationMember.application_id == EventApplication.id)
        .join(Event, EventApplication.event_id == Event.id)
        .where(EventApplicationMember.user_id == user.id)
        .options(
            selectinload(EventApplicationMember.application).selectinload(EventApplication.event),
            selectinload(EventApplicationMember.application).selectinload(EventApplication.band),
        )
        .order_by(Event.event_date.desc(), EventApplication.submitted_at.desc())
    )
    rows = []
    for member in result.scalars().unique():
        application = member.application
        event = application.event
        rows.append(
            {
                "id": application.id,
                "source": "admin",
                "source_label": "админы",
                "event_id": event.id,
                "event_title": event.title,
                "event_date": event.event_date.isoformat(),
                "event_location": event.location,
                "event_status": event.status,
                "band_id": application.band_id,
                "band_name": application.band.name,
                "song_title": application.song_title,
                "instrument_role": member.instrument_role,
                "submitted_at": application.submitted_at.isoformat(),
            }
        )
    personal_result = await session.execute(
        select(PersonalEventHistory)
        .where(PersonalEventHistory.user_id == user.id)
        .order_by(PersonalEventHistory.event_date.desc(), PersonalEventHistory.created_at.desc())
    )
    rows.extend(_personal_event_to_row(item) for item in personal_result.scalars())
    return sorted(rows, key=lambda item: item["event_date"], reverse=True)


async def create_personal_event_history(session: AsyncSession, telegram_id: int, **data) -> PersonalEventHistory:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    item = PersonalEventHistory(user_id=user.id, **data)
    session.add(item)
    await session.flush()
    await log_action(
        session,
        user.id,
        "event.history.personal.create",
        "personal_event_history",
        item.id,
        f"Добавил личное мероприятие: {item.title} ({item.event_date.isoformat()})",
    )
    await session.commit()
    await session.refresh(item)
    return item


async def get_personal_event_history(session: AsyncSession, telegram_id: int, item_id: int) -> dict:
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    result = await session.execute(
        select(PersonalEventHistory).where(PersonalEventHistory.id == item_id, PersonalEventHistory.user_id == user.id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise NotFoundError("Personal event not found")
    return _personal_event_to_row(item)


async def update_personal_event_history(session: AsyncSession, telegram_id: int, item_id: int, **data) -> PersonalEventHistory:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    result = await session.execute(
        select(PersonalEventHistory).where(PersonalEventHistory.id == item_id, PersonalEventHistory.user_id == user.id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise NotFoundError("Personal event not found")
    for key, value in data.items():
        setattr(item, key, value)
    await log_action(
        session,
        user.id,
        "event.history.personal.update",
        "personal_event_history",
        item.id,
        f"Обновил личное мероприятие #{item.id}: {item.title}",
    )
    await session.commit()
    await session.refresh(item)
    return item


async def delete_personal_event_history(session: AsyncSession, telegram_id: int, item_id: int) -> dict:
    from app.services.audit import log_action

    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_approved(user)
    result = await session.execute(
        select(PersonalEventHistory).where(PersonalEventHistory.id == item_id, PersonalEventHistory.user_id == user.id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise NotFoundError("Personal event not found")
    deleted = _personal_event_to_row(item)
    await log_action(
        session,
        user.id,
        "event.history.personal.delete",
        "personal_event_history",
        item.id,
        f"Удалил личное мероприятие #{item.id}: {item.title}",
    )
    await session.delete(item)
    await session.commit()
    return deleted


async def export_event_applications(session: AsyncSession, telegram_id: int, event_id: int) -> bytes:
    user = await get_user_by_telegram_id(session, telegram_id)
    if user is None:
        raise NotFoundError("User not found")
    require_admin(user)
    result = await session.execute(
        select(EventApplication)
        .where(EventApplication.event_id == event_id)
        .options(
            selectinload(EventApplication.event),
            selectinload(EventApplication.band),
            selectinload(EventApplication.members).selectinload(EventApplicationMember.user),
        )
    )
    applications = list(result.scalars())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Participants"
    sheet.append(["Concert", "Band", "Song", "Last name", "First name", "Study group", "Role", "Telegram"])
    first_data_row = sheet.max_row + 1
    for application in applications:
        for member in application.members:
            user = member.user
            sheet.append(
                [
                    application.event.title,
                    application.band.name,
                    application.song_title,
                    user.last_name,
                    user.first_name,
                    user.study_group,
                    member.instrument_role,
                    f"@{user.telegram_username}" if user.telegram_username else "",
                ]
            )
    # User-entered text must never be evaluated as a formula by Excel / LibreOffice.
    for row in sheet.iter_rows(min_row=first_data_row):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
